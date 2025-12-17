"""
xlwings trader report generator.

Reads positions from Excel, calculates risk metrics, and writes
comprehensive report back to Excel.

Usage:
    # From Excel: Click "Generate Report" button
    # From command line: python xlwings_report.py --trader 1
"""

import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from pathlib import Path
import logging
import argparse
from typing import Optional, Dict, Tuple

# Import our modules
from config import (
    TRADER_FILES,
    SHEET_POSITIONS,
    SHEET_SETTINGS,
    SHEET_REPORT,
    DEFAULT_CORRELATION,
    USE_LIVE_CORRELATIONS,
    REQUIRE_LIVE_CORRELATIONS,
    VOL_LOOKBACK_DAYS,
    ZSCORE_LOOKBACK_DAYS,
    CORRELATION_LOOKBACK_DAYS,
    get_data_provider
)
from risk_calcs import (
    calculate_position_vol,
    calculate_portfolio_vol_uniform,
    calculate_marginal_vol,
    calculate_2sigma_drawdowns,
    calculate_entry_score,
    calculate_pnl,
    calculate_distance_to_tp,
    calculate_distance_to_stop,
    calculate_sharpe_ratio
)
from svb_stress_data import load_svb_scenario

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Try to import xlwings (optional for testing)
try:
    import xlwings as xw
    XLWINGS_AVAILABLE = True
except ImportError:
    logger.warning("xlwings not available. Some features will be limited.")
    XLWINGS_AVAILABLE = False


# =============================================================================
# DATA READING
# =============================================================================

def read_positions_openpyxl(filepath: Path) -> pd.DataFrame:
    """
    Read positions from Excel using openpyxl (without xlwings).

    Args:
        filepath: Path to Excel file

    Returns:
        DataFrame with position data
    """
    import openpyxl

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[SHEET_POSITIONS]

    # Read data
    data = []
    headers = [cell.value for cell in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        # Skip empty rows
        if not any(row):
            continue

        data.append(row)

    wb.close()

    df = pd.DataFrame(data, columns=headers)

    # Clean up column names
    df.columns = [col.lower().replace(' ', '_').replace('(', '').replace(')', '').replace(':', '') for col in df.columns]
    
    # Rename common columns to standard names
    column_renames = {
        'bpv_$k': 'bpv',
        'p&l_$': 'pnl',
    }
    df = df.rename(columns=column_renames)
    
    # Calculate total_score from individual score columns if it's None
    # (Excel formulas aren't evaluated when reading with openpyxl data_only=True)
    score_cols = [c for c in df.columns if c.startswith('score_') and c != 'total_score']
    if score_cols and 'total_score' in df.columns:
        # Sum the score columns, treating None/NaN as 0
        df['total_score'] = df[score_cols].apply(
            lambda row: sum(float(v) if pd.notna(v) else 0 for v in row), 
            axis=1
        )
        logger.info(f"Calculated total_score from {len(score_cols)} score columns")

    return df


def read_settings_openpyxl(filepath: Path) -> Dict:
    """
    Read settings from Excel using openpyxl.

    Args:
        filepath: Path to Excel file

    Returns:
        Dictionary with settings
    """
    import openpyxl

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[SHEET_SETTINGS]

    settings = {
        'trader_name': ws['B1'].value or "Unknown Trader",
        'risk_limit': float(ws['B2'].value or 10) * 1_000_000,  # Convert to $
        'default_correlation': float(ws['B3'].value or DEFAULT_CORRELATION),
        'excel_file_path': ws['B4'].value or str(filepath),
    }

    wb.close()

    return settings


def read_positions_xlwings(wb) -> pd.DataFrame:
    """
    Read positions from Excel using xlwings.

    Args:
        wb: xlwings Workbook object

    Returns:
        DataFrame with position data
    """
    ws = wb.sheets[SHEET_POSITIONS]

    # Read all data
    data = ws.range('A1').expand().value

    # Create DataFrame
    df = pd.DataFrame(data[1:], columns=data[0])

    # Clean up column names
    df.columns = [col.lower().replace(' ', '_').replace('(', '').replace(')', '').replace(':', '') for col in df.columns]

    return df


def read_settings_xlwings(wb) -> Dict:
    """Read settings from Excel using xlwings."""
    ws = wb.sheets[SHEET_SETTINGS]

    settings = {
        'trader_name': ws.range('B1').value or "Unknown Trader",
        'risk_limit': float(ws.range('B2').value or 10) * 1_000_000,
        'default_correlation': float(ws.range('B3').value or DEFAULT_CORRELATION),
        'excel_file_path': ws.range('B4').value or "",
    }

    return settings


# =============================================================================
# DATA FILTERING
# =============================================================================

def filter_positions(df: pd.DataFrame, status: str) -> pd.DataFrame:
    """
    Filter positions by status.

    Args:
        df: Positions DataFrame
        status: Status to filter ("Open", "Closed", "Potential")

    Returns:
        Filtered DataFrame
    """
    return df[df['status'].str.upper() == status.upper()].copy()


def get_open_positions(df: pd.DataFrame) -> pd.DataFrame:
    """Get open positions."""
    return filter_positions(df, "Open")


def get_closed_positions(df: pd.DataFrame) -> pd.DataFrame:
    """Get closed positions."""
    return filter_positions(df, "Closed")


def get_potential_positions(df: pd.DataFrame) -> pd.DataFrame:
    """Get potential positions."""
    return filter_positions(df, "Potential")


# =============================================================================
# REPORT GENERATION
# =============================================================================

def enrich_positions_with_market_data(
    positions: pd.DataFrame,
    data_provider
) -> pd.DataFrame:
    """
    Enrich positions with live market data and calculations.

    Args:
        positions: Positions DataFrame
        data_provider: Data provider instance

    Returns:
        Enhanced DataFrame with market data and risk metrics
    """
    df = positions.copy()

    # Add market data columns
    df['live_price'] = None
    df['annual_vol_bp'] = None
    df['zscore'] = None
    df['position_vol'] = None

    for idx, row in df.iterrows():
        ticker = row.get('ticker')
        if not ticker or pd.isna(ticker):
            continue

        try:
            # Get live price
            live_price = data_provider.get_live_price(ticker)
            if live_price is not None:
                df.at[idx, 'live_price'] = live_price

                # Update current_level if missing
                if pd.isna(row.get('current_level')):
                    df.at[idx, 'current_level'] = live_price

            # Get volatility
            annual_vol = data_provider.calculate_annual_vol(ticker, VOL_LOOKBACK_DAYS)
            if annual_vol is not None:
                df.at[idx, 'annual_vol_bp'] = annual_vol

                # Calculate position vol (use abs(bpv) since BPV can be negative for shorts)
                bpv = row.get('bpv', 0)
                if not pd.isna(bpv) and bpv != 0:
                    pos_vol = calculate_position_vol(abs(bpv), annual_vol)
                    df.at[idx, 'position_vol'] = pos_vol

            # Get z-score
            zscore = data_provider.calculate_zscore(ticker, ZSCORE_LOOKBACK_DAYS)
            if zscore is not None:
                df.at[idx, 'zscore'] = zscore

                # Calculate entry score
                direction = row.get('direction', '')
                if direction:
                    entry_score = calculate_entry_score(zscore, direction)
                    df.at[idx, 'score_7_entry_z-score'] = entry_score

        except Exception as e:
            logger.warning(f"Error enriching {ticker}: {e}")

    return df


def generate_current_portfolio_section(
    open_positions: pd.DataFrame,
    portfolio_vol: float,
    correlation: float
) -> pd.DataFrame:
    """
    Generate current portfolio section of report.

    Args:
        open_positions: Open positions DataFrame (enriched with market data)
        portfolio_vol: Current portfolio volatility
        correlation: Correlation assumption

    Returns:
        DataFrame with portfolio summary
    """
    df = open_positions.copy()

    # Calculate P&L for each position (now with product_type support)
    df['pnl'] = df.apply(
        lambda row: calculate_pnl(
            row.get('entry_level', 0),
            row.get('current_level', 0),
            row.get('bpv', 0),
            row.get('direction', ''),
            row.get('product', None)  # Pass product type
        ) if not pd.isna(row.get('entry_level')) else 0,
        axis=1
    )

    # Calculate marginal vol for each position
    df['marginal_vol'] = 0.0
    for idx, row in df.iterrows():
        pos_vol = row.get('position_vol', 0)
        # Handle None/NaN values
        if pos_vol is None or pd.isna(pos_vol):
            pos_vol = 0
        if pos_vol > 0:
            # Remove this position from portfolio vol
            other_vols = df.loc[df.index != idx, 'position_vol'].dropna()
            if len(other_vols) > 0:
                other_port_vol = calculate_portfolio_vol_uniform(other_vols.values, correlation)
            else:
                other_port_vol = 0

            # Marginal vol is the contribution of this position
            marginal = calculate_marginal_vol(pos_vol, other_port_vol, correlation)
            df.at[idx, 'marginal_vol'] = marginal

    # Calculate distance to TP and Stop
    df['distance_to_tp'] = df.apply(
        lambda row: calculate_distance_to_tp(
            row.get('current_level', 0),
            row.get('tp_level', 0),
            row.get('direction', '')
        ) if not pd.isna(row.get('tp_level')) else None,
        axis=1
    )

    df['distance_to_stop'] = df.apply(
        lambda row: calculate_distance_to_stop(
            row.get('current_level', 0),
            row.get('stop_level', 0),
            row.get('direction', '')
        ) if not pd.isna(row.get('stop_level')) else None,
        axis=1
    )

    # Flag positions near stop (within 20%)
    df['near_stop_warning'] = df.apply(
        lambda row: (
            abs(row.get('distance_to_stop', 999)) < 0.2 * abs(row.get('entry_level', 1) - row.get('stop_level', 0)) * 100
            if not pd.isna(row.get('stop_level'))
            else False
        ),
        axis=1
    )

    # Select relevant columns for report
    report_cols = [
        'ticker', 'headline', 'direction', 'bpv', 'entry_level', 'current_level',
        'pnl', 'position_vol', 'marginal_vol', 'tp_level', 'stop_level',
        'distance_to_tp', 'distance_to_stop', 'near_stop_warning'
    ]

    return df[report_cols]


def generate_risk_metrics(
    open_positions: pd.DataFrame,
    settings: Dict,
    portfolio_vol: float
) -> Dict:
    """
    Generate risk metrics section.

    Args:
        open_positions: Open positions DataFrame (enriched)
        settings: Settings dictionary
        portfolio_vol: Portfolio volatility

    Returns:
        Dictionary with risk metrics
    """
    # Portfolio metrics
    daily_vol = portfolio_vol / np.sqrt(252)

    # Calculate 2-sigma drawdowns
    drawdowns = calculate_2sigma_drawdowns(portfolio_vol)

    # Risk utilization
    risk_limit = settings['risk_limit']
    risk_util = (portfolio_vol / risk_limit) * 100 if risk_limit > 0 else 0
    headroom = risk_limit - portfolio_vol

    # Expected return (placeholder - would need Sharpe estimates per position)
    # For now, assume 0.5 Sharpe for all positions
    expected_return = portfolio_vol * 0.5

    # Portfolio Sharpe
    port_sharpe = expected_return / portfolio_vol if portfolio_vol > 0 else 0

    metrics = {
        'portfolio_vol_annual': portfolio_vol,
        'portfolio_vol_daily': daily_vol,
        'expected_return': expected_return,
        'portfolio_sharpe': port_sharpe,
        'risk_limit': risk_limit,
        'risk_utilisation_pct': risk_util,
        'headroom': headroom,
        'drawdown_1d': drawdowns['1d'],
        'drawdown_1w': drawdowns['1w'],
        'drawdown_2w': drawdowns['2w'],
        'drawdown_1m': drawdowns['1m'],
    }

    return metrics


def generate_svb_stress_section(open_positions: pd.DataFrame) -> Tuple[pd.DataFrame, Dict]:
    """
    Generate SVB stress scenario section.

    Args:
        open_positions: Open positions DataFrame

    Returns:
        Tuple of (stress DataFrame, summary dict)
    """
    scenario = load_svb_scenario()

    # Calculate stress for all positions
    stress_df = scenario.calculate_portfolio_stress_pnl(open_positions)

    # Get summary
    summary = scenario.get_portfolio_stress_summary(open_positions)

    return stress_df, summary


def generate_performance_metrics(
    open_positions: pd.DataFrame,
    closed_positions: pd.DataFrame
) -> Dict:
    """
    Generate performance metrics for open and closed trades.

    Args:
        open_positions: Open positions DataFrame (with P&L calculated)
        closed_positions: Closed positions DataFrame

    Returns:
        Dictionary with performance metrics
    """
    # Open positions performance
    total_unrealised_pnl = open_positions['pnl'].sum() if 'pnl' in open_positions.columns else 0

    # Closed positions performance
    closed_total_pnl = 0
    win_rate = 0
    avg_win = 0
    avg_loss = 0
    win_loss_ratio = 0

    if not closed_positions.empty and 'pnl' in closed_positions.columns:
        closed_total_pnl = closed_positions['pnl'].sum()

        winners = closed_positions[closed_positions['pnl'] > 0]
        losers = closed_positions[closed_positions['pnl'] < 0]

        total_closed = len(closed_positions)
        num_winners = len(winners)

        if total_closed > 0:
            win_rate = (num_winners / total_closed) * 100

        if num_winners > 0:
            avg_win = winners['pnl'].mean()

        if len(losers) > 0:
            avg_loss = losers['pnl'].mean()

        if avg_loss != 0:
            win_loss_ratio = abs(avg_win / avg_loss)

    metrics = {
        'total_unrealised_pnl': total_unrealised_pnl,
        'total_realised_pnl': closed_total_pnl,
        'win_rate_pct': win_rate,
        'avg_win': avg_win,
        'avg_loss': avg_loss,
        'win_loss_ratio': win_loss_ratio,
    }

    return metrics


def generate_score_analysis(closed_positions: pd.DataFrame) -> pd.DataFrame:
    """
    Analyze performance by score range.

    Args:
        closed_positions: Closed positions DataFrame

    Returns:
        DataFrame with score analysis
    """
    if closed_positions.empty or 'total_score' not in closed_positions.columns:
        return pd.DataFrame()

    # Define score bins
    bins = [0, 3, 5, 7, 9]
    labels = ['0-3', '3.5-5', '5.5-7', '7.5-9']

    df = closed_positions.copy()
    
    # Filter out rows with None/NaN total_score
    df = df[df['total_score'].notna()]
    
    if df.empty:
        return pd.DataFrame()
    
    df['score_range'] = pd.cut(df['total_score'], bins=bins, labels=labels, include_lowest=True)

    # Group by score range
    analysis = df.groupby('score_range', observed=True).agg({
        'trade_id': 'count',
        'pnl': ['mean', lambda x: (x > 0).sum() / len(x) * 100 if len(x) > 0 else 0]
    })

    analysis.columns = ['num_trades', 'avg_pnl', 'win_rate_pct']

    return analysis.reset_index()


def generate_potential_trade_analysis(
    potential_positions: pd.DataFrame,
    open_positions: pd.DataFrame,
    settings: Dict,
    current_portfolio_vol: float
) -> pd.DataFrame:
    """
    Analyze potential trades and their impact on risk budget.

    Args:
        potential_positions: Potential positions DataFrame (enriched)
        open_positions: Open positions DataFrame
        settings: Settings dictionary
        current_portfolio_vol: Current portfolio volatility

    Returns:
        DataFrame with potential trade analysis
    """
    if potential_positions.empty:
        return pd.DataFrame()

    df = potential_positions.copy()
    correlation = settings['default_correlation']
    risk_limit = settings['risk_limit']
    headroom = risk_limit - current_portfolio_vol

    # For each potential trade
    results = []
    for _, row in df.iterrows():
        trade_vol = row.get('position_vol', 0)

        if trade_vol is not None and trade_vol > 0:
            # Calculate marginal vol
            marginal = calculate_marginal_vol(trade_vol, current_portfolio_vol, correlation)

            # New portfolio vol
            new_port_vol = current_portfolio_vol + marginal

            # New utilization
            new_util = (new_port_vol / risk_limit) * 100 if risk_limit > 0 else 0

            # Fits budget?
            fits = new_port_vol <= risk_limit

            # Calculate max BPV
            from risk_calcs import calculate_max_bpv
            annual_vol_bp = row.get('annual_vol_bp', 0) or 0
            if annual_vol_bp > 0:
                max_bpv = calculate_max_bpv(headroom, current_portfolio_vol, annual_vol_bp, correlation)
            else:
                max_bpv = 0

            results.append({
                'headline': row.get('headline'),
                'ticker': row.get('ticker'),
                'direction': row.get('direction'),
                'proposed_bpv': row.get('bpv'),
                'entry_level': row.get('entry_level'),
                'entry_zscore': row.get('zscore'),
                'trade_vol': trade_vol,
                'marginal_vol': marginal,
                'new_portfolio_vol': new_port_vol,
                'new_risk_util_pct': new_util,
                'fits_budget': fits,
                'max_bpv': max_bpv,
            })

    return pd.DataFrame(results)


def generate_full_report(
    filepath: Path,
    use_xlwings: bool = False
) -> Dict:
    """
    Generate full trader report.

    Args:
        filepath: Path to trader Excel file
        use_xlwings: Whether to use xlwings (True) or openpyxl (False)

    Returns:
        Dictionary with all report sections
    """
    logger.info(f"Generating report for: {filepath}")

    # Read data
    if use_xlwings and XLWINGS_AVAILABLE:
        wb = xw.Book(filepath)
        positions = read_positions_xlwings(wb)
        settings = read_settings_xlwings(wb)
    else:
        positions = read_positions_openpyxl(filepath)
        settings = read_settings_openpyxl(filepath)

    logger.info(f"Read {len(positions)} positions for {settings['trader_name']}")

    # Filter by status
    open_pos = get_open_positions(positions)
    closed_pos = get_closed_positions(positions)
    potential_pos = get_potential_positions(positions)

    logger.info(f"  Open: {len(open_pos)}, Closed: {len(closed_pos)}, Potential: {len(potential_pos)}")

    # Get data provider
    try:
        data_provider = get_data_provider()
        logger.info(f"Using data provider: {data_provider.get_provider_name()}")
    except Exception as e:
        logger.error(f"Could not initialize data provider: {e}")
        logger.warning("Report will have limited functionality without market data")
        data_provider = None

    # Enrich with market data
    if data_provider:
        try:
            open_pos = enrich_positions_with_market_data(open_pos, data_provider)
            potential_pos = enrich_positions_with_market_data(potential_pos, data_provider)
            logger.info("Enriched positions with market data")
        except Exception as e:
            logger.error(f"Error enriching positions: {e}")

    # Calculate portfolio vol (with live correlations if enabled)
    corr_matrix = None
    using_live_correlations = False
    
    if not open_pos.empty and 'position_vol' in open_pos.columns:
        vols = open_pos['position_vol'].dropna()
        if len(vols) > 0:
            # Try to get live correlation matrix if enabled
            if USE_LIVE_CORRELATIONS and data_provider and 'ticker' in open_pos.columns:
                try:
                    tickers = open_pos['ticker'].dropna().unique().tolist()
                    if len(tickers) > 1:
                        corr_matrix = data_provider.get_correlation_matrix(tickers, CORRELATION_LOOKBACK_DAYS)
                        if corr_matrix is not None and not corr_matrix.empty:
                            using_live_correlations = True
                            logger.info(f"✓ Using live correlation matrix ({CORRELATION_LOOKBACK_DAYS} day lookback)")
                        else:
                            error_msg = f"⚠ WARNING: Could not retrieve correlation matrix, falling back to DEFAULT_CORRELATION={DEFAULT_CORRELATION}"
                            logger.warning(error_msg)
                            if REQUIRE_LIVE_CORRELATIONS:
                                raise ValueError(
                                    "REQUIRE_LIVE_CORRELATIONS is True but correlation matrix unavailable. "
                                    "Check data provider connection or disable REQUIRE_LIVE_CORRELATIONS."
                                )
                except Exception as e:
                    error_msg = f"⚠ WARNING: Error getting correlation matrix: {e}"
                    logger.error(error_msg)
                    if REQUIRE_LIVE_CORRELATIONS:
                        raise ValueError(
                            f"REQUIRE_LIVE_CORRELATIONS is True but failed to get correlations: {e}"
                        )
                    logger.warning(f"⚠ Falling back to DEFAULT_CORRELATION={DEFAULT_CORRELATION}")
            
            # Calculate portfolio vol (with matrix or uniform correlation)
            from risk_calcs import calculate_portfolio_vol
            portfolio_vol = calculate_portfolio_vol(
                open_pos, 
                correlation=settings['default_correlation'],
                corr_matrix=corr_matrix
            )
        else:
            portfolio_vol = 0
    else:
        portfolio_vol = 0

    correlation_method = "Live Market Data" if using_live_correlations else f"Fixed ({settings['default_correlation']})"
    logger.info(f"Portfolio vol: ${portfolio_vol:,.0f} | Correlation method: {correlation_method}")

    # Generate report sections
    report = {
        'settings': settings,
        'timestamp': datetime.now(),
        'portfolio_vol': portfolio_vol,
        'correlation_method': correlation_method,
        'using_live_correlations': using_live_correlations,
    }

    # Section 1: Current Portfolio
    if not open_pos.empty:
        report['current_portfolio'] = generate_current_portfolio_section(
            open_pos, portfolio_vol, settings['default_correlation']
        )
    else:
        report['current_portfolio'] = pd.DataFrame()

    # Section 2: Risk Metrics
    report['risk_metrics'] = generate_risk_metrics(open_pos, settings, portfolio_vol)

    # Section 3: SVB Stress
    if not open_pos.empty:
        stress_df, stress_summary = generate_svb_stress_section(open_pos)
        report['svb_stress'] = stress_df
        report['svb_stress_summary'] = stress_summary
    else:
        report['svb_stress'] = pd.DataFrame()
        report['svb_stress_summary'] = {}

    # Section 4: Performance
    report['performance'] = generate_performance_metrics(open_pos, closed_pos)

    # Section 5: Score Analysis
    if not closed_pos.empty:
        report['score_analysis'] = generate_score_analysis(closed_pos)
    else:
        report['score_analysis'] = pd.DataFrame()

    # Section 6: Potential Trades
    if not potential_pos.empty:
        report['potential_trades'] = generate_potential_trade_analysis(
            potential_pos, open_pos, settings, portfolio_vol
        )
    else:
        report['potential_trades'] = pd.DataFrame()

    logger.info("Report generated successfully")

    return report


# =============================================================================
# EXCEL BUTTON INTEGRATION (xlwings)
# =============================================================================

def generate_report_from_excel():
    """
    Generate report from Excel button click (xlwings integration).

    This function is called when the trader clicks the "Generate Report"
    button in Excel. It detects the current workbook and generates the report.

    Usage from VBA:
        RunPython "from xlwings_report import generate_report_from_excel; generate_report_from_excel()"
    """
    if not XLWINGS_AVAILABLE:
        raise ImportError("xlwings is required for Excel button integration")

    try:
        # Get the active workbook (the one the button was clicked in)
        wb = xw.Book.caller()

        logger.info(f"Generating report for: {wb.name}")

        # Read positions and settings
        positions = read_positions_xlwings(wb)
        settings = read_settings_xlwings(wb)

        # Filter by status
        open_pos = get_open_positions(positions)
        closed_pos = get_closed_positions(positions)
        potential_pos = get_potential_positions(positions)

        # Get data provider (if available)
        try:
            data_provider = get_data_provider()
        except Exception as e:
            logger.warning(f"Data provider not available: {e}")
            data_provider = None

        # Enrich with market data (if available)
        if data_provider:
            try:
                from xlwings_report import enrich_positions_with_market_data
                open_pos = enrich_positions_with_market_data(open_pos, data_provider)
                potential_pos = enrich_positions_with_market_data(potential_pos, data_provider)
            except Exception as e:
                logger.warning(f"Could not enrich with market data: {e}")

        # Calculate portfolio vol (with live correlations if enabled)
        corr_matrix = None
        using_live_correlations = False
        
        if not open_pos.empty and 'position_vol' in open_pos.columns:
            vols = open_pos['position_vol'].dropna()
            if len(vols) > 0:
                # Try to get live correlation matrix if enabled
                if USE_LIVE_CORRELATIONS and data_provider and 'ticker' in open_pos.columns:
                    try:
                        tickers = open_pos['ticker'].dropna().unique().tolist()
                        if len(tickers) > 1:
                            corr_matrix = data_provider.get_correlation_matrix(tickers, CORRELATION_LOOKBACK_DAYS)
                            if corr_matrix is not None and not corr_matrix.empty:
                                using_live_correlations = True
                                logger.info(f"✓ Using live correlation matrix ({CORRELATION_LOOKBACK_DAYS} day lookback)")
                            else:
                                logger.warning(f"⚠ WARNING: Could not retrieve correlation matrix, falling back to DEFAULT_CORRELATION={DEFAULT_CORRELATION}")
                                if REQUIRE_LIVE_CORRELATIONS:
                                    raise ValueError("REQUIRE_LIVE_CORRELATIONS is True but correlation matrix unavailable")
                    except Exception as e:
                        logger.error(f"⚠ WARNING: Error getting correlation matrix: {e}")
                        if REQUIRE_LIVE_CORRELATIONS:
                            raise ValueError(f"REQUIRE_LIVE_CORRELATIONS is True but failed to get correlations: {e}")
                        logger.warning(f"⚠ Falling back to DEFAULT_CORRELATION={DEFAULT_CORRELATION}")
                
                # Calculate portfolio vol
                from risk_calcs import calculate_portfolio_vol
                portfolio_vol = calculate_portfolio_vol(
                    open_pos, 
                    correlation=settings['default_correlation'],
                    corr_matrix=corr_matrix
                )
            else:
                portfolio_vol = 0
        else:
            portfolio_vol = 0

        # Write report to Report sheet
        ws = wb.sheets[SHEET_REPORT]
        ws.clear()

        # Header
        ws.range('A1').value = f"TRADER RISK REPORT - {settings['trader_name']}"
        ws.range('A1').font.bold = True
        ws.range('A1').font.size = 16

        ws.range('A2').value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        # Section 1: Portfolio Summary
        row = 4
        ws.range(f'A{row}').value = "PORTFOLIO SUMMARY"
        ws.range(f'A{row}').font.bold = True
        row += 1

        ws.range(f'A{row}').value = "Portfolio Vol:"
        ws.range(f'B{row}').value = portfolio_vol
        ws.range(f'B{row}').number_format = '$#,##0'
        row += 1

        risk_limit = settings.get('risk_limit', 10_000_000)
        risk_util = (portfolio_vol / risk_limit * 100) if risk_limit > 0 else 0
        ws.range(f'A{row}').value = "Risk Utilization:"
        ws.range(f'B{row}').value = risk_util / 100
        ws.range(f'B{row}').number_format = '0.0%'
        row += 1

        ws.range(f'A{row}').value = "Headroom:"
        ws.range(f'B{row}').value = risk_limit - portfolio_vol
        ws.range(f'B{row}').number_format = '$#,##0'
        row += 2

        # Section 2: 2-Sigma Drawdowns
        ws.range(f'A{row}').value = "2-SIGMA DRAWDOWNS"
        ws.range(f'A{row}').font.bold = True
        row += 1

        drawdowns = calculate_2sigma_drawdowns(portfolio_vol)
        for horizon, label in [('1d', '1 Day'), ('1w', '1 Week'), ('2w', '2 Weeks'), ('1m', '1 Month')]:
            ws.range(f'A{row}').value = label
            ws.range(f'B{row}').value = drawdowns[horizon]
            ws.range(f'B{row}').number_format = '$#,##0'
            row += 1

        row += 1

        # Section 3: Open Positions
        ws.range(f'A{row}').value = "OPEN POSITIONS"
        ws.range(f'A{row}').font.bold = True
        row += 1

        if not open_pos.empty:
            # Write headers
            headers = ['Ticker', 'Headline', 'Direction', 'BPV', 'Entry', 'Current', 'P&L', 'Position Vol', 'Score']
            for i, header in enumerate(headers):
                ws.range(row, i+1).value = header
                ws.range(row, i+1).font.bold = True
            row += 1

            # Write positions
            for _, pos in open_pos.iterrows():
                ws.range(row, 1).value = pos.get('ticker', '')
                ws.range(row, 2).value = pos.get('headline', '')
                ws.range(row, 3).value = pos.get('direction', '')
                ws.range(row, 4).value = pos.get('bpv', 0)
                ws.range(row, 5).value = pos.get('entry_level', 0)
                ws.range(row, 6).value = pos.get('current_level', 0)
                ws.range(row, 7).value = pos.get('pnl', 0)
                ws.range(row, 7).number_format = '$#,##0'
                ws.range(row, 8).value = pos.get('position_vol', 0)
                ws.range(row, 8).number_format = '$#,##0'
                ws.range(row, 9).value = pos.get('total_score', 0)
                row += 1

        row += 1

        # Section 4: SVB Stress
        ws.range(f'A{row}').value = "SVB STRESS SCENARIO"
        ws.range(f'A{row}').font.bold = True
        row += 1

        if not open_pos.empty:
            scenario = load_svb_scenario()
            stress_summary = scenario.get_portfolio_stress_summary(open_pos)

            ws.range(f'A{row}').value = "Total Stress P&L:"
            ws.range(f'B{row}').value = stress_summary.get('total_stress_pnl', 0)
            ws.range(f'B{row}').number_format = '$#,##0'
            row += 1

            ws.range(f'A{row}').value = "Worst Position:"
            ws.range(f'B{row}').value = stress_summary.get('worst_position', 0)
            ws.range(f'B{row}').number_format = '$#,##0'
            ws.range(f'C{row}').value = stress_summary.get('worst_position_ticker', '')

        # Auto-fit columns
        ws.autofit()

        logger.info("Report written to Excel successfully")

    except Exception as e:
        logger.error(f"Error generating report from Excel: {e}", exc_info=True)
        if XLWINGS_AVAILABLE:
            import xlwings as xw
            xw.Book.caller().sheets[SHEET_REPORT].range('A1').value = f"ERROR: {str(e)}"


# =============================================================================
# MAIN ENTRY POINT
# =============================================================================

def main():
    """Main entry point for command line usage."""
    parser = argparse.ArgumentParser(description="Generate trader risk report")
    parser.add_argument('--trader', type=int, required=True, help="Trader number (1-5)")
    parser.add_argument('--output', type=str, help="Output file path (optional)")
    parser.add_argument('--xlwings', action='store_true', help="Use xlwings instead of openpyxl")

    args = parser.parse_args()

    # Get trader file
    if args.trader not in TRADER_FILES:
        logger.error(f"Invalid trader number: {args.trader}. Must be 1-5.")
        return

    filepath = TRADER_FILES[args.trader]

    if not filepath.exists():
        logger.error(f"Trader file not found: {filepath}")
        return

    # Generate report
    try:
        report = generate_full_report(filepath, use_xlwings=args.xlwings)

        # Print summary
        print("\n" + "="*60)
        print(f"TRADER REPORT: {report['settings']['trader_name']}")
        print(f"Generated: {report['timestamp'].strftime('%Y-%m-%d %H:%M:%S')}")
        print("="*60)

        print(f"\nPortfolio Vol: ${report['portfolio_vol']:,.0f}")

        metrics = report['risk_metrics']
        print(f"Risk Limit: ${metrics['risk_limit']:,.0f}")
        print(f"Risk Utilisation: {metrics['risk_utilisation_pct']:.1f}%")
        print(f"Headroom: ${metrics['headroom']:,.0f}")

        print(f"\n2-Sigma Drawdowns:")
        print(f"  1-day:   ${metrics['drawdown_1d']:,.0f}")
        print(f"  1-week:  ${metrics['drawdown_1w']:,.0f}")
        print(f"  2-week:  ${metrics['drawdown_2w']:,.0f}")
        print(f"  1-month: ${metrics['drawdown_1m']:,.0f}")

        perf = report['performance']
        print(f"\nPerformance:")
        print(f"  Unrealised P&L: ${perf['total_unrealised_pnl']:,.0f}")
        print(f"  Realised P&L: ${perf['total_realised_pnl']:,.0f}")
        print(f"  Win Rate: {perf['win_rate_pct']:.1f}%")

        if 'svb_stress_summary' in report and report['svb_stress_summary']:
            svb = report['svb_stress_summary']
            print(f"\nSVB Stress Scenario:")
            print(f"  Total Stress P&L: ${svb.get('total_stress_pnl', 0):,.0f}")

        print("\n" + "="*60)

    except Exception as e:
        logger.error(f"Error generating report: {e}", exc_info=True)
        return


if __name__ == "__main__":
    main()
