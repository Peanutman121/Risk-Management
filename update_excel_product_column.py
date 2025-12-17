"""
Script to update Excel files with new Product column structure.

Changes:
1. Add 'Product' column after 'Ticker' column
2. Change Direction column to use Pay/Rec for swaps, Long/Short for futures
3. Infer Product type from tickers and update accordingly
"""

import openpyxl
from pathlib import Path
import re

# Excel files to update
EXCEL_DIR = Path(__file__).parent.parent / "excel"
EXCEL_FILES = [
    EXCEL_DIR / "positions_trader1.xlsx",
    EXCEL_DIR / "positions_trader2.xlsx",
    EXCEL_DIR / "positions_trader3.xlsx",
    EXCEL_DIR / "positions_trader4.xlsx",
    EXCEL_DIR / "positions_trader5.xlsx",
    EXCEL_DIR / "template.xlsx",
]

# Product type inference rules based on ticker patterns
def infer_product_type(ticker: str) -> str:
    """Infer product type from ticker string."""
    if not ticker:
        return "IR Swap"
    
    ticker_upper = ticker.upper()
    
    # Futures patterns (e.g., SFRM5, ERH6, FVH5, TYH5, etc.)
    futures_patterns = [
        r'^SFR[A-Z]\d',   # SOFR futures
        r'^ER[A-Z]\d',    # Euribor futures
        r'^FF[A-Z]\d',    # Fed Funds futures
        r'^ED[A-Z]\d',    # Eurodollar futures
        r'^FV[A-Z]\d',    # 5Y Treasury futures
        r'^TY[A-Z]\d',    # 10Y Treasury futures
        r'^US[A-Z]\d',    # 30Y Treasury futures
        r'^TU[A-Z]\d',    # 2Y Treasury futures
        r'^WN[A-Z]\d',    # Ultra bond futures
        r'^UXY[A-Z]\d',   # Ultra 10Y futures
        r'^RX[A-Z]\d',    # Bund futures
        r'^OE[A-Z]\d',    # Bobl futures
        r'^DU[A-Z]\d',    # Schatz futures
        r'^IK[A-Z]\d',    # BTP futures
        r'^OAT[A-Z]\d',   # OAT futures
        r'^G [A-Z]\d',    # Gilt futures
        r'COMDTY$',       # Any Comdty suffix (futures)
    ]
    
    for pattern in futures_patterns:
        if re.search(pattern, ticker_upper):
            return "Futures"
    
    # XCCY patterns
    if 'XCCY' in ticker_upper or 'XCS' in ticker_upper:
        return "XCCY"
    
    # ASW patterns
    if 'ASW' in ticker_upper:
        return "ASW"
    
    # Basis patterns (e.g., 3s6s, basis swaps)
    if 'BASIS' in ticker_upper or '3S6S' in ticker_upper or '1S3S' in ticker_upper:
        return "Basis"
    
    # Default to IR Swap
    return "IR Swap"


def convert_direction(old_direction: str, product_type: str) -> str:
    """Convert old Long/Short direction to new convention based on product type."""
    if not old_direction:
        return ""
    
    old_upper = old_direction.upper().strip()
    
    if product_type == "Futures":
        # Futures keep Long/Short
        if old_upper == "LONG":
            return "Long"
        elif old_upper == "SHORT":
            return "Short"
    else:
        # Swaps use Pay/Rec
        # Old convention: Long = want rates down (receiver)
        # New convention: Pay = want rates up, Rec = want rates down
        if old_upper == "LONG":
            return "Rec"  # Long position in rate terms = receiver
        elif old_upper == "SHORT":
            return "Pay"  # Short position in rate terms = payer
    
    return old_direction  # Return as-is if not recognized


def update_excel_file(filepath: Path) -> dict:
    """
    Update a single Excel file with Product column.
    
    Returns dict with stats about changes made.
    """
    if not filepath.exists():
        return {"status": "skipped", "reason": "file not found"}
    
    print(f"\nProcessing: {filepath.name}")
    
    wb = openpyxl.load_workbook(filepath)
    ws = wb["Positions"]
    
    # Get current headers
    headers = [cell.value for cell in ws[1]]
    print(f"  Current headers: {headers[:10]}...")
    
    # Check if Product column already exists
    if "Product" in headers:
        print("  Product column already exists - updating values only")
        product_col = headers.index("Product") + 1
        direction_col = headers.index("Direction") + 1 if "Direction" in headers else None
        ticker_col = headers.index("Ticker") + 1 if "Ticker" in headers else None
    else:
        # Find Ticker column to insert Product after it
        ticker_col = headers.index("Ticker") + 1 if "Ticker" in headers else 4
        direction_col = headers.index("Direction") + 1 if "Direction" in headers else None
        
        # Insert new Product column after Ticker
        product_insert_col = ticker_col + 1
        ws.insert_cols(product_insert_col)
        ws.cell(row=1, column=product_insert_col, value="Product")
        
        # Update column indices after insert
        product_col = product_insert_col
        if direction_col and direction_col >= product_insert_col:
            direction_col += 1
        
        print(f"  Inserted Product column at position {product_col}")
    
    # Refresh headers after potential insert
    headers = [cell.value for cell in ws[1]]
    
    # Find column indices
    ticker_col = headers.index("Ticker") + 1 if "Ticker" in headers else None
    product_col = headers.index("Product") + 1 if "Product" in headers else None
    direction_col = headers.index("Direction") + 1 if "Direction" in headers else None
    
    if not all([ticker_col, product_col, direction_col]):
        print(f"  ERROR: Missing required columns. ticker={ticker_col}, product={product_col}, direction={direction_col}")
        return {"status": "error", "reason": "missing columns"}
    
    # Update data rows
    changes = {"products_set": 0, "directions_converted": 0, "rows_processed": 0}
    
    for row in range(2, ws.max_row + 1):
        ticker = ws.cell(row=row, column=ticker_col).value
        if not ticker:
            continue
        
        changes["rows_processed"] += 1
        
        # Infer and set product type
        product_type = infer_product_type(ticker)
        ws.cell(row=row, column=product_col, value=product_type)
        changes["products_set"] += 1
        
        # Convert direction
        old_direction = ws.cell(row=row, column=direction_col).value
        if old_direction:
            new_direction = convert_direction(old_direction, product_type)
            if new_direction != old_direction:
                ws.cell(row=row, column=direction_col, value=new_direction)
                changes["directions_converted"] += 1
                print(f"    Row {row}: {ticker} -> {product_type}, {old_direction} -> {new_direction}")
    
    # Save
    wb.save(filepath)
    wb.close()
    
    print(f"  ✓ Updated {changes['products_set']} products, converted {changes['directions_converted']} directions")
    
    return {"status": "success", **changes}


def main():
    print("=" * 60)
    print("Updating Excel files with Product column")
    print("=" * 60)
    
    results = {}
    for filepath in EXCEL_FILES:
        results[filepath.name] = update_excel_file(filepath)
    
    print("\n" + "=" * 60)
    print("Summary:")
    print("=" * 60)
    for filename, result in results.items():
        print(f"  {filename}: {result['status']}")


if __name__ == "__main__":
    main()
